"""Bounded local file operations, backups and simple document conversion."""
import csv
import io
import json
import math
import os
import shutil
import statistics
import tempfile
import uuid
import zipfile
from pathlib import Path, PureWindowsPath
from xml.sax.saxutils import escape


class Files:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def path(self, name):
        path = (self.root / name).resolve()
        if not path.is_relative_to(self.root) or any(p.casefold()=='.lulu-backups' for p in path.relative_to(self.root).parts):
            raise ValueError('只能操作工作文件夹中的文件，备份通过恢复入口读取。')
        if any(PureWindowsPath(p).is_reserved() or ':' in p for p in path.relative_to(self.root).parts):
            raise ValueError('文件名包含Windows保留名称或不支持的字符')
        return path

    def check_size(self, path):
        if path.stat().st_size > 20 * 1024 * 1024:
            raise ValueError('试用版单个文件上限20MB')
        if path.suffix.lower() in ('.docx', '.xlsx'):
            with zipfile.ZipFile(path) as archive:
                if sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
                    raise ValueError('文档解压大小超过80MB')

    def listing(self, query=''):
        result = []
        for base, dirs, names in os.walk(self.root, followlinks=False):
            dirs[:] = [d for d in dirs if not d.startswith('.') and not (Path(base)/d).is_symlink()]
            for name in names:
                path = Path(base)/name
                if path.is_symlink() or name.startswith('.'):
                    continue
                rel = str(path.relative_to(self.root))
                if query.lower() in rel.lower():
                    result.append({'path': rel, 'bytes': path.stat().st_size})
                if len(result) >= 200:
                    return result
        return result

    def read(self, name):
        path = self.path(name)
        self.check_size(path)
        ext = path.suffix.lower()
        if ext == '.docx':
            from docx import Document
            doc = Document(path)
            return '\n'.join([p.text for p in doc.paragraphs] +
                             ['\t'.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        if ext == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if len(reader.pages) > 150:
                raise ValueError('试用版PDF最多150页')
            return '\n'.join(p.extract_text() or '' for p in reader.pages)
        if ext == '.xlsx':
            return '\n'.join('\t'.join(str(c) if c is not None else '' for c in r)
                             for r in self.table(name))
        if ext not in ('.txt','.md','.csv','.tsv','.json','.html','.log'):
            raise ValueError('暂不支持读取此格式')
        return path.read_text(encoding='utf-8-sig')

    def table(self, name):
        path = self.path(name)
        self.check_size(path)
        if path.suffix.lower() == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = []
                for row in book.active.iter_rows(values_only=True):
                    if len(rows) >= 10000 or len(row) > 100:
                        raise ValueError('试用版表格上限10000行、100列')
                    rows.append(list(row))
                return rows
            finally:
                book.close()
        if path.suffix.lower() not in ('.csv', '.tsv'):
            raise ValueError('分析支持CSV、TSV或XLSX')
        rows = list(csv.reader(io.StringIO(path.read_text(encoding='utf-8-sig')),
                               delimiter='\t' if path.suffix.lower()=='.tsv' else ','))
        if len(rows) > 10000 or any(len(r)>100 for r in rows):
            raise ValueError('试用版表格上限10000行、100列')
        return rows

    def backup(self, path):
        if not path.exists():
            return None
        folder = self.root / '.lulu-backups'
        folder.mkdir(exist_ok=True)
        bid = uuid.uuid4().hex
        shutil.copy2(path, folder / bid)
        (folder / (bid+'.json')).write_text(json.dumps({'path': str(path.relative_to(self.root))}), encoding='utf-8')
        return bid

    def commit(self, path, writer):
        path.parent.mkdir(parents=True, exist_ok=True)
        fd, temp = tempfile.mkstemp(dir=path.parent, suffix=path.suffix)
        os.close(fd)
        try:
            writer(Path(temp))
            self.check_size(Path(temp))
            backup = self.backup(path)
            os.replace(temp, path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        return {'path': str(path.relative_to(self.root)), 'bytes': path.stat().st_size,
                'backup_id': backup}

    def write(self, name, content='', rows=None, overwrite=False):
        path = self.path(name)
        if path.exists() and not overwrite:
            raise ValueError('目标已存在；修改请使用edit或明确overwrite=true，原件将自动备份')
        if len(content)>500000 or (rows is not None and (len(rows)>10000 or any(len(r)>100 for r in rows))):
            raise ValueError('内容超出试用版上限')
        ext = path.suffix.lower()
        def writer(temp):
            if ext in ('.txt','.md','.html','.log'):
                temp.write_text(content, encoding='utf-8')
            elif ext == '.json':
                temp.write_text(json.dumps(json.loads(content), ensure_ascii=False, indent=2), encoding='utf-8')
            elif ext in ('.csv','.tsv'):
                data = rows if rows is not None else list(csv.reader(io.StringIO(content)))
                with temp.open('w', encoding='utf-8-sig', newline='') as handle:
                    csv.writer(handle, delimiter='\t' if ext=='.tsv' else ',').writerows(data)
            elif ext == '.xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                book = Workbook()
                sheet = book.active
                sheet.title = 'Lulu'
                for row in (rows if rows is not None else list(csv.reader(io.StringIO(content)))):
                    sheet.append(row)
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='426B60')
                sheet.freeze_panes = 'A2'
                for col in sheet.columns:
                    sheet.column_dimensions[col[0].column_letter].width = min(50,max(12,max(len(str(c.value or '')) for c in col)+3))
                book.save(temp)
                book.close()
            elif ext == '.docx':
                from docx import Document
                from docx.oxml import OxmlElement
                from docx.oxml.ns import qn
                doc = Document()
                font = doc.styles['Normal'].font
                font.name = 'Microsoft YaHei'
                props = doc.styles['Normal'].element.get_or_add_rPr()
                fonts = OxmlElement('w:rFonts')
                fonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
                props.append(fonts)
                for line in content.splitlines():
                    if line.startswith('# '):
                        doc.add_heading(line[2:], 1)
                    else:
                        doc.add_paragraph(line)
                doc.save(temp)
            elif ext == '.pdf':
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.cidfonts import UnicodeCIDFont
                from reportlab.lib.styles import ParagraphStyle
                from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
                pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
                style = ParagraphStyle('Chinese',fontName='STSong-Light',fontSize=11,leading=18,wordWrap='CJK')
                story=[]
                for line in content.splitlines():
                    story.extend([Paragraph(escape(line) or ' ',style),Spacer(1,5)])
                SimpleDocTemplate(str(temp)).build(story or [Paragraph(' ',style)])
            else:
                raise ValueError('生成支持TXT、MD、JSON、CSV、TSV、XLSX、DOCX、PDF')
        return self.commit(path, writer)

    def edit(self, name, old, new, replace_all=False):
        if not old:
            raise ValueError('旧文本不能为空')
        path = self.path(name)
        self.check_size(path)
        if path.suffix.lower() == '.docx':
            from docx import Document
            doc = Document(path)
            paragraphs = list(doc.paragraphs) + [p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
            if sum(p.text.count(old) for p in paragraphs) != 1:
                raise ValueError('旧文本必须唯一匹配一个段落；请提供更精确的文本')
            for paragraph in paragraphs:
                if old in paragraph.text:
                    for run in paragraph.runs:
                        if old in run.text:
                            run.text = run.text.replace(old,new,1)
                            break
                    else:
                        raise ValueError('文字横跨不同格式片段；此版本不自动改写以免丢失排版')
            return self.commit(path, lambda temp: doc.save(temp))
        if path.suffix.lower() not in ('.txt','.md','.csv','.tsv','.json','.html','.log'):
            raise ValueError('此格式不支持文本替换；XLSX请使用cell操作')
        text = self.read(name)
        if not text.count(old):
            raise ValueError('文件中找不到指定旧文本，请核对文件内容')
        if text.count(old)!=1 and not replace_all:
            raise ValueError('旧文本必须恰好匹配一次；若用户需要全部替换，请设置replace_all=true')
        return self.write(name,text.replace(old,new,-1 if replace_all else 1),overwrite=True)

    def cell(self, name, address, value):
        from openpyxl import load_workbook
        import re
        path = self.path(name)
        if path.suffix.lower()!='.xlsx' or not re.fullmatch(r'[A-Z]{1,2}[1-9][0-9]{0,3}',address):
            raise ValueError('需要XLSX文件和有效单元格地址，例如B2')
        self.check_size(path)
        book = load_workbook(path)
        try:
            book.active[address] = value
            return self.commit(path,lambda temp:book.save(temp))
        finally:
            book.close()

    def convert(self, source, destination):
        src, dst = self.path(source), self.path(destination)
        tables = ('.csv','.tsv','.xlsx')
        if src.suffix.lower() in tables and dst.suffix.lower() in tables:
            return self.write(destination,rows=self.table(source))
        result = self.write(destination,self.read(source))
        result['note'] = '按提取的文本重新生成；不保留原文档复杂排版、图片或公式。'
        return result

    def restore(self, bid):
        if not isinstance(bid,str) or len(bid)!=32 or any(c not in '0123456789abcdef' for c in bid):
            raise ValueError('无效备份编号')
        folder = self.root/'.lulu-backups'
        meta = json.loads((folder/(bid+'.json')).read_text(encoding='utf-8'))
        target = self.path(meta['path'])
        return self.commit(target,lambda temp:shutil.copy2(folder/bid,temp))

    def analyze(self, name):
        rows = self.table(name)
        if not rows:
            return {'rows':0}
        result = {'rows':len(rows)-1,'columns':{}}
        for index, header in enumerate(rows[0]):
            values, missing = [],0
            for row in rows[1:]:
                val = row[index] if index<len(row) else None
                if val in (None,''):
                    missing+=1
                    continue
                try:
                    number=float(val)
                    if math.isfinite(number):
                        values.append(number)
                except (ValueError,TypeError):
                    pass
            stats={'numeric_count':len(values),'missing':missing}
            if values:
                stats.update(sum=sum(values),mean=statistics.mean(values),min=min(values),max=max(values))
            result['columns'][f'{index+1}:{header}']=stats
        return result
